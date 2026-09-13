import os, re, uuid, time, math
from datetime import datetime, timezone
from pathlib import Path
from flask import Flask, request, jsonify, send_from_directory, make_response
from werkzeug.utils import secure_filename

try:
    import fitz  # PyMuPDF
except Exception:
    fitz = None
try:
    from docx import Document as DocxDocument
except Exception:
    DocxDocument = None
try:
    from pptx import Presentation
except Exception:
    Presentation = None
import sqlite3
import requests

BASE_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIR = BASE_DIR / "frontend"
DB_URL = os.getenv("DATABASE_URL", "").strip()
GROQ_API_KEY = os.getenv("GROQ_API_KEY", "").strip()
GROQ_MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-20b")
MAX_FILE_MB = int(os.getenv("MAX_FILE_MB", "15"))
MAX_CHUNKS_TO_SEND = int(os.getenv("MAX_CHUNKS_TO_SEND", "8"))

app = Flask(__name__, static_folder=str(FRONTEND_DIR), static_url_path="")
app.config["MAX_CONTENT_LENGTH"] = MAX_FILE_MB * 1024 * 1024

ALLOWED_EXTS = {"pdf", "docx", "pptx", "txt", "md"}

# ---------------- DB ----------------

def sqlite_conn():
    path = Path(os.getenv("SQLITE_PATH", str(BASE_DIR / "data.db")))
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    return conn


def is_postgres():
    return DB_URL.startswith("postgres")


def db_init():
    if is_postgres():
        import psycopg2
        conn = psycopg2.connect(DB_URL)
        cur = conn.cursor()
        cur.execute("""
        CREATE TABLE IF NOT EXISTS documents (
          id TEXT PRIMARY KEY,
          session_id TEXT NOT NULL,
          filename TEXT NOT NULL,
          file_type TEXT NOT NULL,
          char_count INTEGER NOT NULL,
          created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
        );
        CREATE TABLE IF NOT EXISTS chunks (
          id TEXT PRIMARY KEY,
          document_id TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
          chunk_index INTEGER NOT NULL,
          text TEXT NOT NULL,
          created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
        );
        CREATE TABLE IF NOT EXISTS chats (
          id TEXT PRIMARY KEY,
          session_id TEXT NOT NULL,
          document_id TEXT,
          title TEXT NOT NULL,
          created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
        );
        CREATE TABLE IF NOT EXISTS messages (
          id TEXT PRIMARY KEY,
          chat_id TEXT NOT NULL REFERENCES chats(id) ON DELETE CASCADE,
          role TEXT NOT NULL,
          content TEXT NOT NULL,
          created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
        );
        CREATE INDEX IF NOT EXISTS idx_docs_session ON documents(session_id);
        CREATE INDEX IF NOT EXISTS idx_chunks_doc ON chunks(document_id);
        CREATE INDEX IF NOT EXISTS idx_chats_session ON chats(session_id);
        CREATE INDEX IF NOT EXISTS idx_messages_chat ON messages(chat_id, created_at);
        """)
        conn.commit(); cur.close(); conn.close()
    else:
        conn = sqlite_conn()
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS documents (
          id TEXT PRIMARY KEY, session_id TEXT NOT NULL, filename TEXT NOT NULL,
          file_type TEXT NOT NULL, char_count INTEGER NOT NULL, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS chunks (
          id TEXT PRIMARY KEY, document_id TEXT NOT NULL, chunk_index INTEGER NOT NULL,
          text TEXT NOT NULL, created_at TEXT NOT NULL,
          FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE
        );
        CREATE TABLE IF NOT EXISTS chats (
          id TEXT PRIMARY KEY, session_id TEXT NOT NULL, document_id TEXT,
          title TEXT NOT NULL, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS messages (
          id TEXT PRIMARY KEY, chat_id TEXT NOT NULL, role TEXT NOT NULL,
          content TEXT NOT NULL, created_at TEXT NOT NULL,
          FOREIGN KEY(chat_id) REFERENCES chats(id) ON DELETE CASCADE
        );
        CREATE INDEX IF NOT EXISTS idx_docs_session ON documents(session_id);
        CREATE INDEX IF NOT EXISTS idx_chunks_doc ON chunks(document_id);
        CREATE INDEX IF NOT EXISTS idx_chats_session ON chats(session_id);
        CREATE INDEX IF NOT EXISTS idx_messages_chat ON messages(chat_id, created_at);
        """)
        conn.commit(); conn.close()


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def pg_execute(sql, params=(), fetch=False, many=False):
    import psycopg2
    conn = psycopg2.connect(DB_URL)
    cur = conn.cursor()
    if many: cur.executemany(sql, params)
    else: cur.execute(sql, params)
    rows = cur.fetchall() if fetch else None
    conn.commit(); cur.close(); conn.close()
    return rows


def insert_document(session_id, filename, file_type, char_count):
    did = str(uuid.uuid4()); created = now_iso()
    if is_postgres():
        pg_execute("INSERT INTO documents(id,session_id,filename,file_type,char_count) VALUES(%s,%s,%s,%s,%s)",
                   (did, session_id, filename, file_type, char_count))
    else:
        c = sqlite_conn(); c.execute("INSERT INTO documents VALUES(?,?,?,?,?,?)", (did,session_id,filename,file_type,char_count,created)); c.commit(); c.close()
    return did


def insert_chunks(document_id, chunks):
    if is_postgres():
        pg_execute("INSERT INTO chunks(id,document_id,chunk_index,text) VALUES(%s,%s,%s,%s)",
                   [(str(uuid.uuid4()),document_id,i,t) for i,t in enumerate(chunks)], many=True)
    else:
        c=sqlite_conn(); c.executemany("INSERT INTO chunks VALUES(?,?,?,?,?)", [(str(uuid.uuid4()),document_id,i,t,now_iso()) for i,t in enumerate(chunks)]); c.commit(); c.close()


def get_document(session_id, document_id):
    if is_postgres():
        rows=pg_execute("SELECT id,session_id,filename,file_type,char_count,created_at FROM documents WHERE id=%s AND session_id=%s",(document_id,session_id),True)
        return dict(zip(["id","session_id","filename","file_type","char_count","created_at"], rows[0])) if rows else None
    c=sqlite_conn(); r=c.execute("SELECT * FROM documents WHERE id=? AND session_id=?",(document_id,session_id)).fetchone(); c.close(); return dict(r) if r else None


def get_chunks(document_id):
    if is_postgres():
        rows=pg_execute("SELECT id,chunk_index,text FROM chunks WHERE document_id=%s ORDER BY chunk_index",(document_id,),True)
        return [{"id":r[0],"chunk_index":r[1],"text":r[2]} for r in rows]
    c=sqlite_conn(); rows=c.execute("SELECT id,chunk_index,text FROM chunks WHERE document_id=? ORDER BY chunk_index",(document_id,)).fetchall(); c.close(); return [dict(r) for r in rows]


def list_documents(session_id):
    if is_postgres():
        rows=pg_execute("SELECT id,filename,file_type,char_count,created_at FROM documents WHERE session_id=%s ORDER BY created_at DESC",(session_id,),True)
        return [{"id":r[0],"filename":r[1],"file_type":r[2],"char_count":r[3],"created_at":str(r[4])} for r in rows]
    c=sqlite_conn(); rows=c.execute("SELECT id,filename,file_type,char_count,created_at FROM documents WHERE session_id=? ORDER BY created_at DESC",(session_id,)).fetchall(); c.close(); return [dict(r) for r in rows]


def create_chat(session_id, document_id, title):
    cid=str(uuid.uuid4()); created=now_iso()
    if is_postgres(): pg_execute("INSERT INTO chats(id,session_id,document_id,title) VALUES(%s,%s,%s,%s)",(cid,session_id,document_id,title))
    else:
        c=sqlite_conn(); c.execute("INSERT INTO chats VALUES(?,?,?,?,?)",(cid,session_id,document_id,title,created)); c.commit(); c.close()
    return cid


def add_message(chat_id, role, content):
    mid=str(uuid.uuid4()); created=now_iso()
    if is_postgres(): pg_execute("INSERT INTO messages(id,chat_id,role,content) VALUES(%s,%s,%s,%s)",(mid,chat_id,role,content))
    else:
        c=sqlite_conn(); c.execute("INSERT INTO messages VALUES(?,?,?,?,?)",(mid,chat_id,role,content,created)); c.commit(); c.close()
    return mid


def get_chat(session_id, chat_id):
    if is_postgres():
        rows=pg_execute("SELECT id,session_id,document_id,title,created_at FROM chats WHERE id=%s AND session_id=%s",(chat_id,session_id),True)
        if not rows:return None
        r=rows[0]; chat={"id":r[0],"session_id":r[1],"document_id":r[2],"title":r[3],"created_at":str(r[4])}
        msgs=pg_execute("SELECT role,content,created_at FROM messages WHERE chat_id=%s ORDER BY created_at",(chat_id,),True)
        chat["messages"]=[{"role":m[0],"content":m[1],"created_at":str(m[2])} for m in msgs]
        return chat
    c=sqlite_conn(); r=c.execute("SELECT * FROM chats WHERE id=? AND session_id=?",(chat_id,session_id)).fetchone()
    if not r:c.close();return None
    chat=dict(r); msgs=c.execute("SELECT role,content,created_at FROM messages WHERE chat_id=? ORDER BY created_at",(chat_id,)).fetchall(); c.close(); chat["messages"]=[dict(m) for m in msgs]; return chat


def list_chats(session_id):
    if is_postgres():
        rows=pg_execute("SELECT id,document_id,title,created_at FROM chats WHERE session_id=%s ORDER BY created_at DESC",(session_id,),True)
        return [{"id":r[0],"document_id":r[1],"title":r[2],"created_at":str(r[3])} for r in rows]
    c=sqlite_conn(); rows=c.execute("SELECT id,document_id,title,created_at FROM chats WHERE session_id=? ORDER BY created_at DESC",(session_id,)).fetchall(); c.close(); return [dict(r) for r in rows]


def ensure_session():
    sid=request.cookies.get("rag_session")
    if not sid: sid=str(uuid.uuid4())
    return sid

# ---------------- parsing/chunking/retrieval ----------------

def clean_text(text):
    text=text.replace("\x00", " ").replace("\r\n","\n").replace("\r","\n")
    text=re.sub(r"[ \t]+"," ",text)
    text=re.sub(r"\n{3,}","\n\n",text)
    return text.strip()


def parse_file(file_storage):
    filename=secure_filename(file_storage.filename or "document")
    ext=Path(filename).suffix.lower().lstrip('.')
    if ext not in ALLOWED_EXTS: raise ValueError("Unsupported file type. Use PDF, DOCX, PPTX, TXT, or MD.")
    data=file_storage.read()
    if not data: raise ValueError("The uploaded file is empty.")
    if ext in {"txt","md"}:
        text=data.decode("utf-8",errors="ignore")
    elif ext=="pdf":
        if fitz is None: raise RuntimeError("PyMuPDF is not installed.")
        doc=fitz.open(stream=data,filetype="pdf")
        text="\n\n".join(page.get_text("text") for page in doc)
        doc.close()
    elif ext=="docx":
        if DocxDocument is None: raise RuntimeError("python-docx is not installed.")
        import io
        d=DocxDocument(io.BytesIO(data))
        parts=[p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text="\n".join(parts)
    elif ext=="pptx":
        if Presentation is None: raise RuntimeError("python-pptx is not installed.")
        import io
        prs=Presentation(io.BytesIO(data)); parts=[]
        for i,slide in enumerate(prs.slides,1):
            slide_parts=[f"Slide {i}"]
            for shape in slide.shapes:
                if hasattr(shape,"text") and shape.text.strip(): slide_parts.append(shape.text.strip())
            parts.append("\n".join(slide_parts))
        text="\n\n".join(parts)
    return filename, ext, clean_text(text)


def chunk_text(text, target=1400, overlap=220):
    if len(text)<=target:return [text]
    paras=[p.strip() for p in re.split(r"\n\s*\n",text) if p.strip()]
    chunks=[]; current=""
    for p in paras:
        if len(p)>target:
            sentences=re.split(r"(?<=[.!?])\s+",p)
            for s in sentences:
                if not current: current=s
                elif len(current)+1+len(s)<=target: current += " "+s
                else:
                    chunks.append(current); tail=current[-overlap:] if overlap else ""; current=(tail+" "+s).strip()
            continue
        if not current: current=p
        elif len(current)+2+len(p)<=target: current += "\n\n"+p
        else:
            chunks.append(current); tail=current[-overlap:] if overlap else ""; current=(tail+"\n\n"+p).strip()
    if current:chunks.append(current)
    return chunks


def tokenize(s):
    return [w.lower() for w in re.findall(r"[a-zA-Z0-9_]{2,}",s)]


def score_chunk(query, text):
    q=tokenize(query); t=tokenize(text)
    if not q or not t:return 0.0
    tf={w:t.count(w) for w in set(t)}
    uniq=set(q)
    score=sum(min(tf.get(w,0),3) for w in uniq)
    # small phrase/heading bonus
    qlow=query.lower().strip()
    tlow=text.lower()
    if qlow and qlow in tlow: score += 6
    return score/(math.sqrt(len(t))+1e-6)


def retrieve(query, chunks, k=MAX_CHUNKS_TO_SEND):
    ranked=sorted(chunks,key=lambda c:score_chunk(query,c["text"]),reverse=True)
    return ranked[:k]

# ---------------- Groq ----------------

def groq_chat(messages, temperature=0.15, max_tokens=800):
    if not GROQ_API_KEY:
        raise RuntimeError("GROQ_API_KEY is not configured on the server.")
    url="https://api.groq.com/openai/v1/chat/completions"
    payload={"model":GROQ_MODEL,"messages":messages,"temperature":temperature,"max_tokens":max_tokens}
    r=requests.post(url,headers={"Authorization":f"Bearer {GROQ_API_KEY}","Content-Type":"application/json"},json=payload,timeout=90)
    if r.status_code>=400:
        try: detail=r.json().get("error",{}).get("message",r.text)
        except Exception: detail=r.text
        raise RuntimeError(f"AI provider error: {detail}")
    data=r.json(); return data["choices"][0]["message"]["content"].strip()


def doc_summary(text):
    # Keep prompts bounded for free/low-resource use. Retrieval is not needed for a global summary.
    sample=text[:18000]
    return groq_chat([
        {"role":"system","content":"You summarize documents accurately. Use only the supplied document text. Do not invent facts. Return a useful study/work summary with: Overview, Key points, Important details, and One-line takeaway."},
        {"role":"user","content":f"Summarize this document text:\n\n{sample}"}
    ],temperature=0.1,max_tokens=900)


def answer_question(question, chunks, history):
    selected=retrieve(question,chunks)
    context="\n\n".join([f"[Source chunk {c['chunk_index']+1}]\n{c['text']}" for c in selected])
    hist=history[-6:]
    msgs=[{"role":"system","content":"You are a document-grounded RAG assistant. Answer using only the provided document context. If the answer is not supported by the context, say that the document does not provide enough information. Be concise but useful. Cite supporting chunks like [Source chunk 2]."}]
    for m in hist: msgs.append({"role":m["role"],"content":m["content"]})
    msgs.append({"role":"user","content":f"Document context:\n{context}\n\nQuestion: {question}"})
    return groq_chat(msgs,temperature=0.15,max_tokens=900), selected

# ---------------- routes ----------------

@app.after_request
def headers(resp):
    resp.headers["X-Content-Type-Options"]="nosniff"
    resp.headers["X-Frame-Options"]="DENY"
    resp.headers["Referrer-Policy"]="strict-origin-when-cross-origin"
    return resp

@app.errorhandler(413)
def too_large(_): return jsonify({"error":f"File is too large. Maximum size is {MAX_FILE_MB} MB."}),413

@app.route("/")
def index(): return send_from_directory(FRONTEND_DIR,"index.html")

@app.route("/assets/<path:path>")
def assets(path): return send_from_directory(FRONTEND_DIR / "assets", path)

@app.get("/api/health")
def health(): return jsonify({"ok":True,"model":GROQ_MODEL,"database":"postgres" if is_postgres() else "sqlite"})

@app.get("/api/documents")
def api_documents():
    sid=ensure_session(); resp=make_response(jsonify({"documents":list_documents(sid)})); resp.set_cookie("rag_session",sid,max_age=60*60*24*365,samesite="Lax"); return resp

@app.post("/api/upload")
def upload():
    sid=ensure_session()
    f=request.files.get("file")
    if not f:return jsonify({"error":"Choose a document first."}),400
    try:
        filename,ext,text=parse_file(f)
        if len(text)<30:return jsonify({"error":"I could not extract enough text from this file. Scanned/image-only PDFs need OCR, which this lightweight version does not include yet."}),422
        chunks=chunk_text(text)
        did=insert_document(sid,filename,ext,len(text)); insert_chunks(did,chunks)
        resp=make_response(jsonify({"document":{"id":did,"filename":filename,"file_type":ext,"char_count":len(text),"chunks":len(chunks)}})); resp.set_cookie("rag_session",sid,max_age=60*60*24*365,samesite="Lax"); return resp
    except ValueError as e:return jsonify({"error":str(e)}),400
    except Exception as e:return jsonify({"error":str(e)}),500

@app.post("/api/documents/<document_id>/summarize")
def summarize(document_id):
    sid=ensure_session(); doc=get_document(sid,document_id)
    if not doc:return jsonify({"error":"Document not found."}),404
    chunks=get_chunks(document_id); text="\n\n".join(c["text"] for c in chunks)
    try:return jsonify({"summary":doc_summary(text)})
    except Exception as e:return jsonify({"error":str(e)}),502

@app.get("/api/chats")
def chats():
    sid=ensure_session(); resp=make_response(jsonify({"chats":list_chats(sid)})); resp.set_cookie("rag_session",sid,max_age=60*60*24*365,samesite="Lax"); return resp

@app.post("/api/chats")
def create_chat_route():
    sid=ensure_session(); data=request.get_json(silent=True) or {}; did=data.get("document_id")
    if did and not get_document(sid,did):return jsonify({"error":"Document not found."}),404
    title=(data.get("title") or "New document chat")[:100]; cid=create_chat(sid,did,title)
    resp=make_response(jsonify({"chat":{"id":cid,"document_id":did,"title":title}})); resp.set_cookie("rag_session",sid,max_age=60*60*24*365,samesite="Lax"); return resp

@app.get("/api/chats/<chat_id>")
def get_chat_route(chat_id):
    sid=ensure_session(); chat=get_chat(sid,chat_id)
    if not chat:return jsonify({"error":"Chat not found."}),404
    return jsonify({"chat":chat})

@app.post("/api/chats/<chat_id>/messages")
def message_route(chat_id):
    sid=ensure_session(); chat=get_chat(sid,chat_id)
    if not chat:return jsonify({"error":"Chat not found."}),404
    data=request.get_json(silent=True) or {}; q=(data.get("message") or "").strip()
    if not q:return jsonify({"error":"Message cannot be empty."}),400
    if not chat.get("document_id"):return jsonify({"error":"Attach a document to this chat first."}),400
    chunks=get_chunks(chat["document_id"])
    try:
        answer, used=answer_question(q,chunks,chat["messages"])
        add_message(chat_id,"user",q); add_message(chat_id,"assistant",answer)
        return jsonify({"answer":answer,"sources":[{"chunk":c["chunk_index"]+1,"preview":c["text"][:280]} for c in used]})
    except Exception as e:return jsonify({"error":str(e)}),502

@app.post("/api/chats/<chat_id>/rename")
def rename_chat(chat_id):
    sid=ensure_session(); chat=get_chat(sid,chat_id)
    if not chat:return jsonify({"error":"Chat not found."}),404
    title=((request.get_json(silent=True) or {}).get("title") or "Document chat")[:100]
    if is_postgres(): pg_execute("UPDATE chats SET title=%s WHERE id=%s",(title,chat_id))
    else: c=sqlite_conn(); c.execute("UPDATE chats SET title=? WHERE id=?",(title,chat_id)); c.commit(); c.close()
    return jsonify({"ok":True,"title":title})

@app.post("/api/chats/<chat_id>/clear")
def clear_chat(chat_id):
    sid=ensure_session(); chat=get_chat(sid,chat_id)
    if not chat:return jsonify({"error":"Chat not found."}),404
    if is_postgres(): pg_execute("DELETE FROM messages WHERE chat_id=%s",(chat_id,))
    else: c=sqlite_conn(); c.execute("DELETE FROM messages WHERE chat_id=?",(chat_id,)); c.commit(); c.close()
    return jsonify({"ok":True})

@app.delete("/api/documents/<document_id>")
def delete_document(document_id):
    sid=ensure_session(); doc=get_document(sid,document_id)
    if not doc:return jsonify({"error":"Document not found."}),404
    if is_postgres():
        pg_execute("DELETE FROM documents WHERE id=%s",(document_id,))
    else:
        c=sqlite_conn(); c.execute("PRAGMA foreign_keys=ON"); c.execute("DELETE FROM documents WHERE id=?",(document_id,)); c.commit(); c.close()
    return jsonify({"ok":True})

if __name__=="__main__":
    db_init()
    port=int(os.getenv("PORT","5000"))
    app.run(host="0.0.0.0",port=port,debug=True)
else:
    db_init()
